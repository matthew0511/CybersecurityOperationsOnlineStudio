"""Local PDF and spreadsheet exports; reports are NOT compliance certifications."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from .service import StudioService


def _safe_excel(s: object) -> str:
    """Prevent spreadsheet formula execution on untrusted text."""
    text = str(s if s is not None else "")
    if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return "'" + text
    return text


def export_pdf(svc: StudioService, org: str, destination: Path) -> Path:
    record = svc.organization(org)
    risks, tasks, assessments = svc.list_risks(org), svc.list_tasks(org), svc.list_assessments(org)
    destination = Path(destination)
    styles = getSampleStyleSheet()
    story = [Paragraph("Cyber Security Operations Online Studio", styles["Title"]),
             Paragraph("Assessment & Risk Snapshot", styles["Heading2"]),
             Paragraph(f"Organization: {escape(record['name'])}", styles["Normal"]),
             Paragraph("Generated: " + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"), styles["Normal"]),
             Spacer(1, 16),
             Paragraph(f"Assessments: {len(assessments)} | Risks: {len(risks)} | Remediation tasks: {len(tasks)}", styles["Normal"]),
             Spacer(1, 12), Paragraph("Risk Register", styles["Heading2"])]
    rows = [["Risk title", "Likelihood", "Impact", "Score", "Band", "Status"]]
    for r in risks:
        rows.append([Paragraph(escape(r['title'][:70]), styles["BodyText"]),
                     str(r['likelihood']), str(r['impact']), str(r['score']), r['band'], r['status']])
    if not risks:
        rows.append(["No risks entered", "", "", "", "", ""])
    table = Table(rows, colWidths=[175, 65, 50, 44, 63, 70], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#193049")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), .35, colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f8fb")]),
    ]))
    story += [table, Spacer(1, 16), Paragraph("Important limitations", styles["Heading2"]),
              Paragraph("This report summarizes manually entered information. It is not a security audit, "
                        "HIPAA certification, regulatory compliance determination, or independent verification "
                        "of control effectiveness. Ratings use the draft 5×5 risk matrix (v1).", styles["BodyText"])]
    doc = SimpleDocTemplate(str(destination), pagesize=letter, leftMargin=42, rightMargin=42)
    doc.build(story)
    return destination


def export_xlsx(svc: StudioService, org: str, destination: Path) -> Path:
    record = svc.organization(org)
    risks, tasks, assessments = svc.list_risks(org), svc.list_tasks(org), svc.list_assessments(org)
    wb = Workbook()
    intro = wb.active
    intro.title = "Overview"
    intro.append(["Cyber Security Operations Online Studio — assessment and risk snapshot"])
    intro.append(["Organization", _safe_excel(record["name"])])
    intro.append(["Generated UTC", datetime.now(timezone.utc).isoformat(timespec="seconds")])
    intro.append(["Assessments", len(assessments)])
    intro.append(["Risks", len(risks)])
    intro.append(["Remediation tasks", len(tasks)])
    intro.append(["Notice", "Manual records; not proof of compliance or verified control effectiveness."])
    definitions = [
        ("Risks", ["Title", "Description", "Likelihood", "Impact", "Score", "Band", "Rationale", "Status"],
         [[r[k] for k in ("title", "description", "likelihood", "impact", "score", "band", "rationale", "status")] for r in risks]),
        ("Remediation", ["Title", "Owner", "Due date", "Priority", "Status", "Risk ID"],
         [[t[k] for k in ("title", "owner", "due_date", "priority", "status", "risk_id")] for t in tasks]),
        ("Assessments", ["Title", "Scope", "Framework", "Status", "Created"],
         [[a[k] for k in ("title", "scope", "framework_label", "status", "created_at")] for a in assessments]),
    ]
    for name, headers, data in definitions:
        ws = wb.create_sheet(name)
        ws.append(headers)
        for line in data:
            ws.append([_safe_excel(v) if isinstance(v, str) else v for v in line])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="193049")
        for column in ws.columns:
            letter_name = column[0].column_letter
            ws.column_dimensions[letter_name].width = min(54, max(16, max(len(str(c.value or "")) for c in column[:100]) + 2))
            for cell in column[1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.save(destination)
    return Path(destination)
