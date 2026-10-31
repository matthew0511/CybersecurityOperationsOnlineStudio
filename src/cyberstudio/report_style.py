"""Shared visual styling for existing management report exports."""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.chart import BarChart, Reference
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import TableStyle

NAVY='112C3A'
TEAL='0C827D'
PALE='F0F6F7'

def pdf_styles():
    styles=getSampleStyleSheet()
    for name in ('Title','Heading1','Heading2'):
        styles[name].textColor=colors.HexColor('#'+NAVY)
    styles['Title'].fontSize=23
    styles['Title'].leading=28
    styles['Heading2'].fontSize=12
    styles['Heading2'].spaceBefore=14
    styles['BodyText'].fontSize=9
    styles['BodyText'].leading=14
    styles['Normal'].fontSize=9
    styles['Normal'].leading=14
    return styles

def pdf_table_style():
    return TableStyle([
        ('BACKGROUND',(0,0),(-1,0),colors.HexColor('#'+NAVY)),
        ('TEXTCOLOR',(0,0),(-1,0),colors.white),
        ('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),
        ('FONTSIZE',(0,0),(-1,-1),9),
        ('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#'+PALE)]),
        ('VALIGN',(0,0),(-1,-1),'TOP'),
        ('LEFTPADDING',(0,0),(-1,-1),12),('RIGHTPADDING',(0,0),(-1,-1),12),
        ('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9),
        ('ALIGN',(-1,1),(-1,-1),'RIGHT'),
        ('LINEBELOW',(0,0),(-1,0),2,colors.HexColor('#'+TEAL)),
    ])

def pdf_page(canvas,doc):
    canvas.saveState()
    width,height=doc.pagesize
    canvas.setFillColor(colors.HexColor('#'+TEAL))
    canvas.rect(0,height-10,width,10,fill=1,stroke=0)
    canvas.setFont('Helvetica-Bold',8)
    canvas.setFillColor(colors.HexColor('#'+NAVY))
    canvas.drawString(46,height-32,'CSOS  /  CYBER SECURITY OPERATIONS ONLINE STUDIO')
    canvas.setStrokeColor(colors.HexColor('#DCE5E8'))
    canvas.line(46,38,width-46,38)
    canvas.setFont('Helvetica',7)
    canvas.drawString(46,25,'Program Owner - Matthew Harris  |  Confidential management records')
    canvas.drawRightString(width-46,25,f'{doc.page}')
    canvas.restoreState()

def style_workbook(wb,summary_header=5):
    """Preserve cell values and sheet schemas; add accessible report formatting."""
    for index,ws in enumerate(wb):
        header=summary_header if index==0 else 1
        ws.sheet_view.showGridLines=False
        ws.sheet_properties.tabColor=TEAL if index==0 else NAVY
        ws.freeze_panes=f'A{header+1}'
        ws.auto_filter.ref=f'A{header}:{ws.cell(ws.max_row,ws.max_column).coordinate}'
        ws.print_title_rows=f'1:{header}'
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        ws.oddFooter.center.text='Program Owner - Matthew Harris | Confidential | Page &P of &N'
        for column in ws.columns:
            width=min(56,max(18,max(len(str(c.value or '')) for c in column[:80])+3))
            ws.column_dimensions[column[0].column_letter].width=width
            for cell in column:
                cell.font=Font(name='Calibri',size=11,color=NAVY)
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                if cell.row>header and cell.row%2==0:cell.fill=PatternFill('solid',fgColor=PALE)
        for cell in ws[header]:
            cell.font=Font(name='Calibri',bold=True,color='FFFFFF',size=11)
            cell.fill=PatternFill('solid',fgColor=NAVY)
        ws.row_dimensions[header].height=28
        for row in range(header+1,ws.max_row+1):ws.row_dimensions[row].height=34
        if index==0:
            for cell in ws[1]:
                cell.font=Font(name='Calibri',size=18,bold=True,color='FFFFFF')
                cell.fill=PatternFill('solid',fgColor=TEAL)
            ws.row_dimensions[1].height=42
            ws.column_dimensions['A'].width=58;ws.column_dimensions['B'].width=65
            for row in range(2,header):ws.row_dimensions[row].height=54 if row==4 else 28
            end=ws.max_row-(1 if ws.cell(ws.max_row,1).value=='Disclaimer' else 0)
            if end>header:
                chart=BarChart();chart.type='bar';chart.style=13
                chart.title='Current management indicators';chart.y_axis.title='Metric';chart.x_axis.title='Count'
                chart.add_data(Reference(ws,min_col=2,min_row=header,max_row=end),titles_from_data=True)
                chart.set_categories(Reference(ws,min_col=1,min_row=header+1,max_row=end))
                chart.height=15;chart.width=25;chart.legend=None
                if chart.series:chart.series[0].graphicalProperties.solidFill=TEAL
                ws.add_chart(chart,'D2')
