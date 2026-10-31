"""Plaintext board-ready governance exports; protect files after creation."""
from __future__ import annotations
from pathlib import Path
from xml.sax.saxutils import escape
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet

DISCLAIMER="Local operator-entered records. Reviewer names are not authenticated electronic signatures. Not a certification or independent audit."

def board_packet(service,org,path):
    path=Path(path)
    snap=service.board_snapshot(org)
    m=snap["governance"];e=snap["enterprise"]
    cards=[
        ("Active frameworks",snap["frameworks_enabled"]),
        ("Verified items",snap["implementation_counts"]["verified"]),
        ("High/critical open findings",m["findings_high_critical"]),
        ("Overdue audit findings",m["findings_overdue"]),
        ("Changes awaiting CAB",m["changes_pending_cab"]),
        ("Rolled-back changes",m["changes_rolled_back"]),
        ("Verified evidence",e["evidence_verified"]),
        ("Expired evidence",e["evidence_expired"]),
        ("Open audits",m["audits_open"]),
    ]
    if path.suffix.lower()==".xlsx":
        wb=Workbook();sheet=wb.active;sheet.title="Board Summary"
        sheet.append(["Cybersecurity Governance | Board Packet",snap["as_of"]])
        sheet.append(["Metric","Value"])
        for key,value in cards:sheet.append([key,value])
        sheet.append(["Disclaimer",DISCLAIMER])
        for cell in sheet[1]+sheet[2]:
            cell.font=Font(bold=True,color="FFFFFF")
            cell.fill=PatternFill("solid",fgColor="203D5B")
        sheet.column_dimensions["A"].width=49;sheet.column_dimensions["B"].width=78
        for label,items,fields in (
            ("Changes",service.changes(org),["title","change_type","risk_level","owner","start_at","end_at","status"]),
            ("Audits",service.audits(org),["title","audit_type","framework","lead","start_date","end_date","status"]),
            ("Findings",service.findings(org),["title","severity","owner","due_date","status"]),
            ("Evidence",service.evidence_inventory(org),["title","classification","sha256","review_state","expires_on"])):
            ws=wb.create_sheet(label);ws.append(fields)
            for obj in items:ws.append([str(obj.get(f,"")) for f in fields])
            ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
            for col in "ABCDEFG":ws.column_dimensions[col].width=28
        wb.save(path)
    elif path.suffix.lower()==".pdf":
        styles=getSampleStyleSheet()
        story=[Paragraph("Cybersecurity Governance | Board Packet",styles["Title"]),
               Paragraph("As of: "+escape(snap["as_of"]),styles["Normal"]),
               Spacer(1,16)]
        grid=[["Governance metric","Current value"]]+[[str(a),str(b)] for a,b in cards]
        t=Table(grid,colWidths=[310,175],repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#203D5B")),
                               ("TEXTCOLOR",(0,0),(-1,0),colors.white),
                               ("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#F1F5FA")]),
                               ("BOTTOMPADDING",(0,0),(-1,-1),9)]))
        story.extend([t,Spacer(1,18),Paragraph("Priority audit findings",styles["Heading2"])])
        for f in service.findings(org):
            if f["severity"] in {"high","critical"} and f["status"] not in {"closed","accepted"}:
                story.append(Paragraph(escape(f["severity"].upper()+" | "+f["title"]+" | "+f["owner"]),styles["BodyText"]))
        story.extend([Spacer(1,18),Paragraph("Management caveat",styles["Heading2"]),
                      Paragraph(DISCLAIMER,styles["BodyText"])])
        SimpleDocTemplate(str(path),pagesize=letter,title="Governance Board Packet").build(story)
    else:
        raise ValueError("Board report requires .pdf or .xlsx")
    return path
