"""Immutable report editions, exact PDF previews, template presets and editable slides."""
import hashlib,io,json,tempfile
from datetime import datetime,timezone
from pathlib import Path
from .store import fail,uid,history
from ..database import service

TEMPLATES={
 'executive':dict(title='Executive overview',sections=['summary','risks','tasks','vendors']),
 'board':dict(title='Board oversight',sections=['summary','risks','tasks','vendors','frameworks','findings','history']),
 'audit':dict(title='Audit and control review',sections=['summary','frameworks','findings','history']),
 'technical':dict(title='Technical operations',sections=['risks','tasks','findings']),
 'client':dict(title='Client program review',sections=['summary','risks','tasks','frameworks']),
}
MIMES={'pdf':'application/pdf','xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','pptx':'application/vnd.openxmlformats-officedocument.presentationml.presentation'}


def powerpoint(payload,brand):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    from PIL import Image
    presentation=Presentation();presentation.slide_width=Inches(13.333);presentation.slide_height=Inches(7.5)
    navy=brand['primary_color'];accent=brand['accent_color']
    from cyberstudio.report_style import ink
    heading=navy if ink(navy)=='FFFFFF' else '112C3A'
    def text(slide,x,y,w,h,value,size=18,bold=False,color='112C3A'):
        box=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h));frame=box.text_frame;frame.word_wrap=True
        for i,line in enumerate(str(value).split('\n')):
            p=frame.paragraphs[0] if i==0 else frame.add_paragraph();p.text=line;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        return box
    def slide(title):
        result=presentation.slides.add_slide(presentation.slide_layouts[6]);band=result.shapes.add_shape(1,0,0,presentation.slide_width,Inches(.12));band.fill.solid();band.fill.fore_color.rgb=RGBColor.from_string(accent);band.line.fill.background()
        left=.6
        if brand.get('logo'):
            image=Image.open(io.BytesIO(brand['logo']));ratio=min(1.8/image.width,.65/image.height)
            result.shapes.add_picture(io.BytesIO(brand['logo']), Inches(.6), Inches(.3),width=Inches(image.width*ratio),height=Inches(image.height*ratio));left=2.65
        text(result,left,.3,10,.5,brand['display_name'],16,True,heading)
        text(result,.6,1.25,12,1,title,28,True,heading)
        text(result,.6,6.85,10,.3,'Program Owner - Matthew Harris | '+brand.get('footer_note','')[:100],9)
        text(result,12,6.85,.7,.3,len(presentation.slides),9)
        return result
    cover=slide('Security program | '+payload.get('report_title','Management report'))
    text(cover,.65,2.6,11.8,1,brand.get('tagline',''),24)
    text(cover,.65,4.05,11.8,1.5,'As of '+payload['as_of']+'\n'+payload['period_note'],14)
    if payload.get('commentary'):
        chunks=[payload['commentary'][i:i+900] for i in range(0,len(payload['commentary']),900)]
        for chunk in chunks:text(slide('Management commentary'),.65,2.4,12,4,chunk,18)
    for name,items in payload['sections'].items():
        for offset in range(0,max(1,len(items)),7):
            current=slide(name.replace('_',' ').title()+(' · continued' if offset else ''))
            if not items:text(current,.65,2.5,12,1,'No records in the selected scope.',18)
            for index,item in enumerate(items[offset:offset+7]):
                line=' | '.join(k.replace('_',' ').title()+': '+str(v) for k,v in item.items())
                text(current,.65,2.3+index*.55,12,.55,line[:260],15)
    out=io.BytesIO();presentation.save(out);return out.getvalue()


def freeze(db,user,org,body):
    from .store import access
    access(db,user,org,False,'reports:export')
    template=body.get('template','board');format=body.get('format','pdf')
    if template not in TEMPLATES or format not in MIMES:fail('Choose a supported report template and format')
    from ..experience import report_payload,render_custom_report
    settings={**body,'sections':body.get('sections') or TEMPLATES[template]['sections']}
    payload=report_payload(db,org,settings);payload['report_title']=TEMPLATES[template]['title'];brand=service(db).report_branding(org)
    if any(len(rows)>2000 for rows in payload['sections'].values()):fail('Narrow the period or selected sections to at most 2000 records per section')
    content=powerpoint(payload,brand) if format=='pptx' else render_custom_report(db,org,payload,format,brand)
    digest=hashlib.sha256(content).hexdigest();ident=uid();manifest={k:v for k,v in brand.items() if k!='logo'};manifest['logo_sha256']=hashlib.sha256(brand['logo']).hexdigest() if brand['logo'] else None
    source={'payload':payload,'branding':manifest,'generated_at':datetime.now(timezone.utc).isoformat()}
    db.execute('INSERT INTO pl_editions VALUES (%s,%s,%s,%s,%s,%s::jsonb,%s,%s,%s,CURRENT_TIMESTAMP)',(ident,org,TEMPLATES[template]['title'],format,template,json.dumps(source,default=str),digest,content,user['username']))
    history(db,org,ident,user['username'],'report:frozen',{'sha256':digest,'format':format,'template':template})
    return {'id':ident,'title':TEMPLATES[template]['title'],'format':format,'sha256':digest,'bytes':len(content),'preview_url':f'/api/platform/{org}/editions/{ident}?inline=true' if format=='pdf' else None,'download_url':f'/api/platform/{org}/editions/{ident}'}


def variance(db,org):
    rows=db.execute('SELECT day,metrics FROM xp_snapshots WHERE organization_id=%s ORDER BY day DESC LIMIT 2',(org,)).fetchall()
    if len(rows)<2:return {'available':False,'message':'Two recorded daily snapshots are required. Earlier exposure is not invented.'}
    latest,previous=rows;changes=[]
    for name,value in latest['metrics'].items():
        before=previous['metrics'].get(name)
        if isinstance(value,(int,float)) and not isinstance(value,bool) and isinstance(before,(int,float)):
            changes.append({'metric':name,'previous':before,'current':value,'change':value-before})
    activity=db.execute('SELECT entity_type,action,recorded_at FROM audit_events WHERE organization_id=%s AND recorded_at>=%s ORDER BY recorded_at DESC LIMIT 100',(org,str(previous['day']))).fetchall()
    return {'available':True,'from':previous['day'],'to':latest['day'],'changes':changes,'supporting_activity':[dict(r) for r in activity],'interpretation':'Activity provides context; it does not prove causation or independent control effectiveness.'}
