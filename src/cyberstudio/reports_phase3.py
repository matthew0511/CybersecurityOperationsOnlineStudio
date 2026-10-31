"""Management reporting with explicit manual/unverified-state disclosures.

Exports are plaintext; caller must protect them as sensitive business artifacts.
"""
from __future__ import annotations
from datetime import datetime, timezone
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether

from .phase3 import Phase3Service
from .reports import _safe_excel
from .report_style import pdf_styles, pdf_table_style, pdf_page, style_workbook

DISCLAIMER = ('Operator-entered and imported records only. Framework mappings and manual reviewer '
              'statements are NOT third-party-verified evidence of compliance. No certification, '
              'independent assurance, regulatory conclusion, or security guarantee is implied.')


def _safe_cell(v: object) -> object:
    return _safe_excel(v) if isinstance(v,str) else v


def export_executive_xlsx(svc: Phase3Service, org: str, destination: Path) -> Path:
    record=svc.organization(org)
    m=svc.executive_metrics(org)
    wb=Workbook();summary=wb.active;summary.title='Executive Overview'
    for row in [ ['Cyber Security Operations Online Studio — Executive snapshot'],
                 ['Organization',_safe_cell(record['name'])],
                 ['Generated UTC',datetime.now(timezone.utc).isoformat(timespec='seconds')],
                 ['Disclosure',DISCLAIMER], ['Metric','Value'] ]:
        summary.append(row)
    for k,v in m.items():summary.append([k.replace('_',' ').title(),v])

    def add(name: str, headers: list[str], rows: list[dict], fields: list[str]):
        ws=wb.create_sheet(name);ws.append(headers)
        for r in rows:
            ws.append([_safe_cell(r.get(k)) for k in fields])

    add('Risks', ['Title','Inherent score','Residual score','Acceptance state','Status','Rationale'],
        svc.list_risks(org),['title','score','residual_score','acceptance_state','status','rationale'])
    add('Remediation',['Title','Owner','Due date','Priority','Status'],svc.list_tasks(org),
        ['title','owner','due_date','priority','status'])
    add('Control Evaluations',['Framework','Code','Title','Implementation','Test','Verified evidence supports pass','Evaluator','Rationale'],
        svc.control_evaluations(org),['framework_key','control_code','control_title','implementation','test_result','verified_pass','evaluator','rationale'])
    add('Vendors',['Name','Criticality','Data access','BAA applicability','Review state','Risk score','Next review'],
        svc.vendors(org),['name','criticality','data_access','baa_applicability','review_state','risk_score','next_review'])
    add('Vulnerabilities',['Title','Severity','CVSS entered','Status','Source','External ID','Asset','Due','Overdue','Notes'],
        svc.vulnerabilities(org),['title','severity','cvss','status','source','external_id','asset_name','due_on','overdue','notes'])
    add('Evidence Inventory',['Title','SHA-256','Original review status','Effective status','Expires','Reviewed by'],
        svc.evidence(org),['title','sha256','review_status','effective_status','expires_on','reviewed_by'])
    add('Imported Controls',['Control name','Source ID','Source status','Verification'],
        svc.imported_control_inventory(org),['title','source_control_id','source_implementation','verification'])
    add('Policies',['Policy','Version','Owner','Next review','Status'],svc.policies(org),
        ['title','version','owner','next_review_on','status'])
    add('Incident Inventory',['Incident','Source ID','Severity','Source status','Unverified'],svc.incidents(org),
        ['title','source_incident_id','severity','source_status','imported_unverified'])

    for ws in wb:
        ws.freeze_panes='A2' if ws.title!='Executive Overview' else 'A6'
        ws.auto_filter.ref=ws.dimensions if ws.title!='Executive Overview' else 'A5:B'+str(ws.max_row)
        if ws.title!='Executive Overview':
            for cell in ws[1]:
                cell.font=Font(bold=True,color='FFFFFF')
                cell.fill=PatternFill('solid',fgColor='193049')
        for column in ws.columns:
            col_letter=column[0].column_letter
            ws.column_dimensions[col_letter].width=min(58,max(16,max(len(str(c.value or '')) for c in column[:80])+2))
            for cell in column:cell.alignment=Alignment(vertical='top',wrap_text=True)
    destination=Path(destination)
    style_workbook(wb,summary_header=5)
    wb.save(destination)
    return destination


def export_executive_pdf(svc: Phase3Service,org: str,destination: Path) -> Path:
    record=svc.organization(org); m=svc.executive_metrics(org)
    styles=pdf_styles()
    story=[Paragraph('Cyber Security Operations Online Studio',styles['Title']),
           Paragraph('Executive operations snapshot',styles['Heading2']),
           Paragraph('Organization: '+escape(record['name']),styles['Normal']),
           Paragraph('Generated UTC: '+datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M'),styles['Normal']),
           Spacer(1,12),Paragraph('Risk & Operations Indicators',styles['Heading2'])]
    metric_labels=[
        ('High inherent risks open','high_inherent_open'),('High residual risks open','high_residual_open'),
        ('Active risk acceptances','acceptances_active'),('Expired risk acceptances','acceptances_expired'),
        ('Controls evaluated','controls_evaluated'),('Controls with operator-entered passing test','controls_test_pass'),
        ('Passing tests with currently verified evidence','controls_evidence_supported'),
        ('Expired evidence artifacts','evidence_expired'),('Unverified or expired evidence','evidence_unverified'),
        ('Vendors with missing/overdue reviews','vendor_reviews_overdue'),
        ('High/critical open vulnerabilities','vulnerabilities_high_open'),
        ('Overdue vulnerabilities','vulnerabilities_overdue'),('Overdue remediation','remediation_overdue')]
    rows=[['Indicator','Count']]+[[label,str(m[key])] for label,key in metric_labels]
    t=Table(rows,colWidths=[380,80],repeatRows=1,hAlign='LEFT')
    t.setStyle(pdf_table_style())
    story += [t,Spacer(1,15),Paragraph('Highest-rated inherent risks (up to 10)',styles['Heading2'])]
    risks=svc.list_risks(org)[:10]
    if not risks:story.append(Paragraph('No risk records entered.',styles['Normal']))
    for r in risks:
        story.append(Paragraph(f"{escape(r['title'][:95])} — score {r['score']} / 25; "
                  f"residual {r['residual_score'] if r['residual_score'] is not None else 'not assessed'}; "
                  f"acceptance {escape(r['acceptance_state'])}", styles['BodyText']))
        story.append(Spacer(1,5))
    story += [Spacer(1,16),Paragraph('Interpretation and limitations',styles['Heading2']),
              Paragraph(escape(DISCLAIMER),styles['BodyText']),
              Spacer(1,8),Paragraph('The control counts describe only individually recorded evaluations; '
                  'their denominator is not the full set of framework requirements. Passing evaluations '
                  'are operator-reported, even where an evidence review exists. Missing evidence and '
                  'expired reviews need follow-up. The source data may be incomplete.',styles['BodyText'])]
    doc=SimpleDocTemplate(str(destination),pagesize=letter,leftMargin=46,rightMargin=46,topMargin=58,bottomMargin=52)
    doc.build(story,onFirstPage=pdf_page,onLaterPages=pdf_page)
    return Path(destination)
