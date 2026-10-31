"""Stateful acceptance and security checks for the new platform."""
import io,json,time,hashlib
from datetime import date,timedelta
import pytest,pyotp,httpx
from fastapi.testclient import TestClient
from test_web import web,login,call,PW
from test_experience import invited,send
from cyberstudio.web.app import app
from cyberstudio.web.database import connect
from cyberstudio.web.platform import catalog,security,quarantine,connectors,reporting,jobs
from cyberstudio.web.platform.schema import actor
from cyberstudio.web.platform.store import validate_rules


def setup(web):
    login(web);return call(web,'create_organization',name='Synthetic Platform'),call(web,'create_organization',name='Isolated Organization')

def post(web,org,path,body,expected=200):
    r=web.post('/api/platform/'+org+path,json=body)
    assert r.status_code==expected,r.text
    return r.json()

def save(web,org,kind,title='Example',data=None,**extra):return post(web,org,'/records/'+kind,dict(title=title,data=data or {},**extra))

def test_catalog_and_rule_cycle_contract():
    assert [f['number'] for f in catalog.features()]==list(range(1,51))
    assert len(catalog.modules())>=20
    with pytest.raises(Exception):validate_rules([{'parent':'A','child':'B'},{'parent':'B','child':'A'}])
    validate_rules([{'parent':'A','child':'B'},{'parent':'B','child':'C'}])


def test_versions_scope_and_typed_references(web):
    org,other=setup(web);site=save(web,org,'site')
    result=save(web,org,'site','Edited',id=site['id'],version=site['version']);assert result['version']==2
    post(web,org,'/records/site',{'id':site['id'],'version':1,'title':'Stale','data':{}},409)
    post(web,other,'/records/asset',{'title':'Wrong scope','data':{'site_id':site['id']}},404)
    post(web,org,'/records/asset',{'title':'Invalid','data':{'dependencies':[{}]}},422)
    child=save(web,org,'site','Child',{'parent_id':site['id']})
    post(web,org,'/records/site',{'id':site['id'],'version':2,'title':'Cycle','data':{'parent_id':child['id']}},422)
    assert len(web.get(f'/api/platform/{org}/record/{site["id"]}').json()['history'])==2


def test_independent_approval_acknowledgment_and_capabilities(web):
    org,_=setup(web);reviewer=invited(web,org,'reviewer','reviewer');viewer=invited(web,org,'viewer','viewer')
    policy=save(web,org,'policy','Policy v1',{'version_label':'1','content':'Synthetic policy','approver':'reviewer'})
    policy=post(web,org,'/record/'+policy['id']+'/status',{'version':1,'status':'submitted'})
    post(web,org,'/record/'+policy['id']+'/status',{'version':2,'status':'approved','password':PW},403)
    policy=post(reviewer,org,'/record/'+policy['id']+'/status',{'version':2,'status':'approved','password':PW})
    assert policy['data']['approved_by']=='reviewer'
    post(viewer,org,'/record/'+policy['id']+'/acknowledge',{'version':policy['version'],'confirmed':True})
    result=web.get(f'/api/platform/{org}/record/{policy["id"]}').json();assert result['acknowledgments'][0]['username']=='viewer'
    post(web,org,'/capabilities',{'username':'reviewer','denied':['reports:export'],'password':PW})
    post(reviewer,org,'/editions',{'template':'board','format':'pdf'},403)


def test_mfa_replay_recovery_and_session_revocation(web,monkeypatch):
    setup(web);clock=int(time.time());monkeypatch.setattr(security.time,'time',lambda:clock)
    begin=web.post('/api/security/mfa_begin',json={'password':PW});assert begin.status_code==200
    secret=begin.json()['secret'];code=pyotp.TOTP(secret).at(clock)
    activated=web.post('/api/security/mfa_enable',json={'code':code});assert activated.status_code==200,activated.text
    codes=activated.json()['recovery_codes'];assert len(codes)==8
    another=TestClient(app,base_url='https://testserver')
    assert another.post('/api/login',json={'username':'owner','password':PW}).status_code==401
    assert another.post('/api/login',json={'username':'owner','password':PW,'mfa_code':code}).status_code==401
    clock+=30;new=pyotp.TOTP(secret).at(clock)
    r=another.post('/api/login',json={'username':'owner','password':PW,'mfa_code':new});assert r.status_code==200,r.text
    assert another.post('/api/login',json={'username':'owner','password':PW,'mfa_code':new}).status_code==401
    r=another.post('/api/login',json={'username':'owner','password':PW,'mfa_code':codes[0]});assert r.status_code==200
    assert another.post('/api/login',json={'username':'owner','password':PW,'mfa_code':codes[0]}).status_code==401
    sessions=web.get('/api/security').json()['sessions'];assert len(sessions)>=2
    web.post('/api/security/revoke_others',json={})
    assert another.get('/api/me').status_code==401


def test_password_recovery_is_one_use_and_keeps_mfa(web):
    setup(web);kit=web.post('/api/security/recovery_kit',json={'password':PW}).json()['recovery_codes']
    code=kit[0];new='Different-synthetic-password-2026!'
    public=TestClient(app,base_url='https://testserver')
    r=public.post('/api/recovery/password',json={'username':'owner','recovery_code':code,'password':new});assert r.status_code==200,r.text
    assert web.get('/api/me').status_code==401
    assert public.post('/api/recovery/password',json={'username':'owner','recovery_code':code,'password':new}).status_code==403
    assert public.post('/api/login',json={'username':'owner','password':PW}).status_code==401
    assert public.post('/api/login',json={'username':'owner','password':new}).status_code==200


def test_mfa_requirement_does_not_lock_out_enrollment(web):
    org,_=setup(web)
    assert web.post('/api/security/mfa_policy',json={'required':True,'password':PW}).status_code==409
    with connect() as db:db.execute("INSERT INTO pl_installation(key,value) VALUES ('security_policy','{\"require_mfa\":true}')")
    assert web.get('/api/me').json()['enrollment_required'] is True
    assert web.post('/api/call/list_risks',json={'arguments':{'org':org}}).status_code==403
    assert web.get('/api/security').status_code==200
    assert web.post('/api/security/mfa_begin',json={'password':PW}).status_code==200


def test_runbooks_bia_and_execution_order(web):
    org,_=setup(web);asset=save(web,org,'asset','Service',{'asset_type':'business_service'})
    bia=save(web,org,'bia','BIA',{'asset_id':asset['id'],'rto_hours':8,'rpo_hours':4});assert bia['data']['rto_hours']==8
    runbook=save(web,org,'runbook','Recovery',{'steps':[{'title':'Restore','instructions':'Use isolated backup'},{'title':'Verify','verification':'Readiness check'}]})
    run=post(web,org,'/record/'+runbook['id']+'/execute',{})['id']
    post(web,org,'/executions/'+run,{'version':1,'index':1,'status':'completed','notes':'Not yet'},409)
    post(web,org,'/executions/'+run,{'version':1,'index':0,'status':'completed','notes':'Restored'})
    post(web,org,'/executions/'+run,{'version':2,'index':1,'status':'completed','notes':'Verified'})
    result=web.get(f'/api/platform/{org}/executions/{run}').json();assert result['status']=='completed' and result['steps'][1]['completed_by']=='owner'


def test_catalog_provenance_scoring_and_questionnaire_rules(web):
    org,_=setup(web);data=b'code,title,parent_code\nA,Example control,\nA.1,Child,A\n'
    base={'name':'Synthetic Framework','edition':'1','source':'Owner supplied','license_note':'Synthetic owned content'}
    url=f'/api/platform/{org}/catalogs'
    assert web.post(url,data=base,files={'file':('catalog.csv',data,'text/csv')}).status_code==422
    r=web.post(url,data={**base,'confirmed':'true'},files={'file':('catalog.csv',data,'text/csv')});assert r.status_code==200,r.text
    catalog_id=r.json()['id'];control=save(web,org,'control','Internal control',{'catalog_mappings':[{'catalog_id':catalog_id,'code':'A.1'}]})
    assessment=call(web,'create_assessment',org=org,title='Review',template_id='builtin:third-party-risk-v1')
    rule=save(web,org,'questionnaire_logic','Logic',{'assessment_id':assessment,'rules':[{'parent':'TPR-001','operator':'equals','value':'yes','child':'TPR-002'}]})
    post(web,org,'/record/'+rule['id']+'/status',{'version':1,'status':'active'})
    questions=web.get(f'/api/experience/{org}/questionnaire/{assessment}').json()['questions'];assert next(q for q in questions if q['code']=='TPR-002')['visible'] is False
    rubric=save(web,org,'rubric','Scoring v1',{'method_version':'1','rules':[{'code':'TPR-001','weight':2,'criteria':'Manual 0–100'},{'code':'TPR-002','weight':2,'criteria':'Manual 0–100'}]})
    score=save(web,org,'assessment_score','Review score',{'assessment_id':assessment,'rubric_id':rubric['id'],'scores':[{'code':'TPR-001','score':80,'rationale':'Synthetic explicit rating'}]})
    assert score['data']['calculated']['weighted_score']==80 and score['data']['calculated']['coverage_percent']==50


def test_quarantine_sensitivity_and_nonclean_verdicts(web,monkeypatch):
    org,_=setup(web)
    content=b'-----BEGIN PRIVATE KEY-----\nsynthetic only'
    result=web.post('/api/upload/evidence',data={'org':org,'title':'Sensitive test'},files={'file':('test.txt',content,'text/plain')})
    assert result.status_code==200,result.text
    pending=result.json()['result'];assert pending['pending'] is True
    ident=pending['quarantine_id']
    post(web,org,'/quarantine/'+ident+'/release',{'password':PW,'reason':'Too short'},409)
    release=post(web,org,'/quarantine/'+ident+'/release',{'password':PW,'reason':'This is a documented synthetic test fixture, not a production secret.'});assert release['evidence_id']
    scanner=quarantine.client('s3');scanner.scan_status='THREATS_FOUND'
    bad=web.post('/api/upload/evidence',data={'org':org,'title':'Malicious test'},files={'file':('test.txt',b'opaque sample','text/plain')}).json()['result']
    post(web,org,'/quarantine/'+bad['quarantine_id']+'/release',{'password':PW,'reason':'Long reason cannot override an independent malicious verdict.'},409)
    assert web.post('/api/upload/evidence',data={'org':org,'title':'EICAR'},files={'file':('test.txt',b'EICAR-STANDARD-ANTIVIRUS-TEST-FILE','text/plain')}).status_code==422


def test_frozen_reports_preview_and_slides(web):
    org,_=setup(web)
    for format in ['pdf','xlsx','pptx']:
        edition=post(web,org,'/editions',{'template':'executive','format':format})
        content=web.get(edition['download_url']).content
        assert hashlib.sha256(content).hexdigest()==edition['sha256']
        if format=='pdf':assert web.get(edition['preview_url']).content==content
        if format=='pptx':
            from pptx import Presentation
            presentation=Presentation(io.BytesIO(content));assert len(presentation.slides)>=2
            assert all(any('Program Owner - Matthew Harris' in s.text for s in slide.shapes if s.has_text_frame) for slide in presentation.slides)
    assert web.get(f'/api/platform/{org}/variance').json()['available'] is False


def test_integrations_encrypted_staging_and_host_validation(web,monkeypatch):
    org,_=setup(web)
    body={'provider':'github','name':'Test GitHub','config':{'owner':'example','repository':'demo'},'credentials':{'token':'synthetic-provider-token'}}
    row=post(web,org,'/integrations',body);assert 'secret_cipher' not in row and 'synthetic-provider-token' not in json.dumps(row)
    post(web,org,'/integrations/'+row['id']+'/enable',{'enabled':True},409)
    post(web,org,'/integrations',{**body,'config':{'owner':'https://127.0.0.1','repository':'demo'}},422)
    original=connectors.Adapter.__init__
    def initialize(self,row,transport=None):return original(self,row,httpx.MockTransport(lambda req:httpx.Response(200,json={'id':123}) if req.url.path.endswith('/demo') else httpx.Response(200,json=[{'id':1,'title':'Synthetic issue'}])))
    monkeypatch.setattr(connectors.Adapter,'__init__',initialize)
    assert post(web,org,'/integrations/'+row['id']+'/test',{})['authenticated'] is True
    post(web,org,'/integrations/'+row['id']+'/enable',{'enabled':True})
    job=post(web,org,'/integrations/'+row['id']+'/pull',{})['job_id']
    with connect() as db,db.transaction():jobs.tick(db)
    with connect() as db:
        assert db.execute('SELECT status FROM pl_jobs WHERE id=%s',(job,)).fetchone()['status']=='succeeded'
        encrypted=db.execute('SELECT secret_cipher FROM pl_integrations WHERE id=%s',(row['id'],)).fetchone()['secret_cipher'];assert 'synthetic-provider-token' not in encrypted
    staged=web.get(f'/api/platform/{org}/integrations/{row["id"]}/staged').json();assert len(staged)==1
    post(web,org,'/integrations/'+row['id']+'/publish',{'title':'Example','message':'Synthetic'},422)


def test_tokens_are_scoped_expiring_and_revocable(web):
    org,other=setup(web);save(web,org,'site')
    issued=post(web,org,'/tokens',{'name':'Read only','scopes':['records:read'],'days':1,'password':PW})
    client=TestClient(app,base_url='https://testserver');client.headers['Authorization']='Bearer '+issued['token']
    assert len(client.get(f'/api/v1/{org}/records/site').json())==1
    assert client.get(f'/api/v1/{other}/records/site').status_code==403
    assert client.post(f'/api/v1/{org}/records/site',json={'title':'Denied','data':{}}).status_code==403
    post(web,org,'/tokens',{'revoke':issued['id'],'password':PW})
    assert client.get(f'/api/v1/{org}/records/site').status_code==403


def test_audit_rooms_selected_scope_one_use_and_revocation(web):
    org,_=setup(web);first=save(web,org,'site','Included office');second=save(web,org,'site','Excluded office')
    room=post(web,org,'/rooms',{'title':'External review','days':1,'allowed':[{'kind':'program','id':first['id']}]})
    external=TestClient(app,base_url='https://testserver')
    redeemed=external.post('/api/audit-room/redeem',json={'token':room['token']});assert redeemed.status_code==200,redeemed.text
    external.headers['X-CSRF-Token']=redeemed.json()['csrf']
    assert external.post('/api/audit-room/redeem',json={'token':room['token']}).status_code==403
    data=external.get('/api/audit-room').json();assert len(data['records'])==1 and second['id'] not in json.dumps(data)
    assert external.get('/api/me').status_code==401
    assert external.get('/api/audit-room/download/evidence/not-selected').status_code==403
    assert external.post('/api/audit-room/questions',json={'body':'Please explain the office recovery scope.'}).status_code==200
    post(web,org,'/rooms',{'revoke':room['id']})
    assert external.get('/api/audit-room').status_code==401


def test_archive_restore_and_demo_portfolio(web):
    org,other=setup(web);record=save(web,org,'asset','Archived asset')
    post(web,org,'/record/'+record['id']+'/archive',{})
    assert web.get(f'/api/platform/{org}/records/asset').json()==[]
    post(web,org,'/trash',{'kind':'program','id':record['id'],'restore':True})
    assert len(web.get(f'/api/platform/{org}/records/asset').json())==1
    viewer=invited(web,org,'portfolio_viewer','viewer')
    assert [r['id'] for r in viewer.get('/api/platform/portfolio').json()]==[org]
    result=web.post('/api/platform/demo',json={'confirmed':True});assert result.status_code==200,result.text
    assert web.get('/api/platform/'+result.json()['organization_id']+'/overview').json()['demo'] is True
    docs=web.get('/api/platform/documentation').json();assert len(docs['features'])==50 and len(docs['sections'])>=10


def test_database_rls_fails_closed_for_non_superuser(web):
    org,other=setup(web);save(web,org,'site','Allowed');save(web,other,'site','Other')
    viewer=invited(web,org,'rls_viewer','viewer')
    import psycopg
    with connect() as db:
        db.execute("DO $$ BEGIN IF NOT EXISTS(SELECT 1 FROM pg_roles WHERE rolname='pl_rls_probe') THEN CREATE ROLE pl_rls_probe NOLOGIN NOBYPASSRLS; END IF; END $$")
        db.execute('GRANT USAGE ON SCHEMA studio TO pl_rls_probe');db.execute('GRANT SELECT,INSERT ON ALL TABLES IN SCHEMA studio TO pl_rls_probe')
        with db.transaction():
            db.execute('SET LOCAL ROLE pl_rls_probe')
            assert db.execute('SELECT count(*) AS n FROM pl_records').fetchone()['n']==0
            actor(db,'rls_viewer',org)
            rows=db.execute('SELECT organization_id,title FROM pl_records').fetchall();assert len(rows)==1 and rows[0]['organization_id']==org
            with pytest.raises(psycopg.errors.InsufficientPrivilege):
                with db.transaction():db.execute("INSERT INTO pl_records(id,organization_id,kind,title,created_by,updated_by) VALUES ('bad-scope',%s,'site','Denied','rls_viewer','rls_viewer')",(other,))
