from __future__ import annotations
from datetime import date, timedelta
from io import StringIO
from pathlib import Path
import csv
import sqlite3
import zipfile

import pytest
from openpyxl import load_workbook
from cyberstudio.phase3 import Phase3Service
from cyberstudio.phase2 import Phase2Service
from cyberstudio.reports_phase3 import export_executive_pdf, export_executive_xlsx
from cyberstudio.service import NotFoundError,ValidationError
from cyberstudio.vault import Vault

PASSWORD='example-long-test-password-123'

@pytest.fixture
def env(tmp_path):
    with Vault(tmp_path/'customer.csvault',PASSWORD,create=True) as vault:
        s=Phase3Service(vault);a=s.create_organization('Alpha');b=s.create_organization('Bravo')
        yield s,a,b,vault

def csv_blob(fields,rows):
    stream=StringIO();writer=csv.DictWriter(stream,fieldnames=fields)
    writer.writeheader();writer.writerows(rows)
    return stream.getvalue().encode('utf-8')

def test_version3_and_upgrade_phase2_vault(tmp_path):
    path=tmp_path/'upgrade.csvault'
    with Vault(path,PASSWORD,create=True) as vault:
        s=Phase2Service(vault);a=s.create_organization('Existing Org')
        risk=s.add_risk(a,'Legacy risk','',4,3,'Legacy assessment')
        evidence=s.add_evidence(a,'Evidence','test.txt',b'Legacy data')
        assert int(vault.db.execute('PRAGMA user_version').fetchone()[0])==2
    with Vault(path,PASSWORD) as vault:
        s=Phase3Service(vault)
        assert int(vault.db.execute('PRAGMA user_version').fetchone()[0])==3
        assert s.list_risks(a)[0]['id']==risk
        assert s.evidence_bytes(a,evidence)==b'Legacy data'
        assert s.evidence(a)[0]['expires_on']==''
        s.set_evidence_expiration(a,evidence,(date.today()+timedelta(days=30)).isoformat())
        v=s.create_vendor(a,'Existing Vendor')
    with Vault(path,PASSWORD) as vault:
        s=Phase3Service(vault)
        assert s.vendors(a)[0]['id']==v
        assert s.evidence(a)[0]['expires_on']!=''


def test_control_evaluations_require_valid_mapping_and_evidence(env):
    s,a,b,vault=env
    e=s.add_evidence(a,'MFA proof','mfa.txt',b'synthetic MFA proof')
    with pytest.raises(ValidationError):s.evaluate_control(a,'NIST-CSF-2.0','BOGUS','implemented','pass','Owner','Evidence',e)
    with pytest.raises(ValidationError):s.evaluate_control(a,'NIST-CSF-2.0','PR.AA','partial','pass','Owner','Evidence',e)
    s.evaluate_control(a,'NIST-CSF-2.0','PR.AA','implemented','pass','Owner','Attested preliminary test',e)
    assert s.control_evaluations(a)[0]['verified_pass'] is False
    s.review_evidence(a,e,'verified','Reviewer','Checked synthetic proof')
    s.evaluate_control(a,'NIST-CSF-2.0','PR.AA','implemented','pass','Owner','Re-tested and linked reviewed evidence',e)
    assert s.control_evaluations(a)[0]['verified_pass'] is True
    s.set_evidence_expiration(a,e,(date.today()-timedelta(days=1)).isoformat())
    assert s.control_evaluations(a)[0]['verified_pass'] is False
    assert s.evidence(a)[0]['effective_status']=='expired'
    assert s.executive_metrics(a)['evidence_expired']==1
    assert s.control_evaluations(b)==[]
    with pytest.raises(NotFoundError):s.evaluate_control(b,'NIST-CSF-2.0','PR.AA','implemented','pass','Owner','Cross-org evidence',e)
    with pytest.raises(sqlite3.IntegrityError):
        with vault.transaction() as conn:conn.execute('DELETE FROM control_evaluations')


def test_vendor_review_audit_expiration_and_isolation(env):
    s,a,b,_=env
    vendor=s.create_vendor(a,'Cloud Provider','SaaS','critical','regulated','yes','Security owner')
    with pytest.raises(ValidationError):s.create_vendor(a,'Cloud Provider')
    with pytest.raises(NotFoundError):s.record_vendor_review(b,vendor,'Bob','Not allowed',3,4,(date.today()+timedelta(days=90)).isoformat())
    with pytest.raises(ValidationError):s.record_vendor_review(a,vendor,'Bob','Insufficient',3,4,date.today().isoformat())
    review=s.record_vendor_review(a,vendor,'Bob','Validated by operator',3,5,(date.today()+timedelta(days=45)).isoformat(),'needs_remediation')
    row=s.vendors(a)[0]
    assert (row['risk_score'],row['review_state'])==(15,'needs_remediation')
    assert len(s.vendor_reviews(a,vendor))==1
    assert s.vendors(b)==[]


def test_vulnerability_history_isolation_and_input_validation(env):
    s,a,b,vault=env
    vuln=s.add_vulnerability(a,'CVE example','manual','CVE-2026-99999','critical',9.8,'test server',
        (date.today()-timedelta(days=3)).isoformat())
    assert s.vulnerabilities(a)[0]['overdue'] is True
    with pytest.raises(ValidationError):s.add_vulnerability(a,'Duplicate','manual','CVE-2026-99999')
    with pytest.raises(NotFoundError):s.update_vulnerability(b,vuln,'mitigated','Bob','Not my tenant')
    with pytest.raises(ValidationError):s.add_vulnerability(a,'Not numeric','manual',cvss=11)
    s.update_vulnerability(a,vuln,'mitigated','Jane','Patch tested')
    assert s.vulnerabilities(a)[0]['overdue'] is False
    assert s.executive_metrics(a)['vulnerabilities_high_open']==0
    with pytest.raises(sqlite3.IntegrityError):
        with vault.transaction() as conn:conn.execute('DELETE FROM vulnerability_history')


def test_phase3_risk_acceptance_requires_vault_reauthentication(env):
    s,a,b,v=env
    rid=s.add_risk(a,'Test risk','',3,4,'Risk basis')
    s.set_residual_risk(a,rid,2,2,'Tested safeguards')
    due=(date.today()+timedelta(days=60)).isoformat()
    with pytest.raises(ValidationError):s.record_risk_decision(a,rid,'approved','Boss','Accepted',due,True)
    with pytest.raises(ValidationError):s.record_risk_decision(a,rid,'approved','Boss','Accepted',due,True,operator_passphrase='incorrect')
    assert s.acceptance_state(a,rid)=='not_approved'
    assert s.record_risk_decision(a,rid,'approved','Boss','Accepted',due,True,operator_passphrase=PASSWORD)
    assert s.acceptance_state(a,rid)=='active'
    assert any(x['action']=='vault_operator_reauthenticated' for x in s.audit_events(a))


def test_notion_multi_import_previews_and_atomic_validation(env):
    s,a,b,_=env
    cases={
        'compliance_controls':(['Control Name','Control ID','Implementation Status','Client'],[
             {'Control Name':'MFA','Control ID':'PR.AA','Implementation Status':'Implemented','Client':'Alpha'}]),
        'security_policies':(['Policy Name','Approval Status','Version'],[
             {'Policy Name':'Access Standard','Approval Status':'Approved','Version':'2.0'}]),
        'vendor_assessments':(['Vendor Name','Assessment Status','Assessment Findings','Criticality Tier','BAA Applicability'],[
             {'Vendor Name':'Example Cloud','Assessment Status':'Approved','Assessment Findings':'Sample gap','Criticality Tier':'High','BAA Applicability':'Required'}]),
        'incident_tracker':(['Incident Name','Status','Severity'],[
             {'Incident Name':'Phishing simulation','Status':'Mitigated','Severity':'Low'}]),
    }
    for kind,(cols,rows) in cases.items():
        payload=csv_blob(cols,rows)
        assert s.preview_notion_module(a,payload,kind)['ready']==1
        with pytest.raises(ValidationError):s.import_notion_module(a,payload,kind,'sample.csv')
        assert s.import_notion_module(a,payload,kind,'sample.csv',confirmed=True)['imported']==1
        with pytest.raises(ValidationError):s.import_notion_module(a,payload,kind,'sample.csv',confirmed=True)
    assert s.imported_control_inventory(a)[0]['verification']=='unverified'
    assert s.policies(a)[0]['status']=='unverified'
    assert s.vendors(a)[0]['baa_applicability']=='unknown'
    assert s.imported_vendor_annotations(a)[0]['source_findings']=='Sample gap'
    assert s.incidents(a)[0]['imported_unverified']==1
    assert s.policies(b)==[] and s.vendors(b)==[]
    bad=csv_blob(['Policy Name','Approval Status','Client'],[
        {'Policy Name':'Good','Approval Status':'Approved','Client':'Alpha'},
        {'Policy Name':'Other Org','Approval Status':'Approved','Client':'Bravo'}])
    assert s.preview_notion_module(a,bad,'security_policies')['skipped']==1
    with pytest.raises(ValidationError):s.import_notion_module(a,bad,'security_policies','bad.csv',confirmed=True)
    assert all(p['title']!='Good' for p in s.policies(a))


def test_reports_include_phase3_and_avoid_formula_injection(env,tmp_path):
    s,a,b,_=env
    s.add_risk(a,'=HYPERLINK("evil")','',4,5,'Risk is synthetic')
    s.create_vendor(a,'=HYPERLINK("evil")')
    s.add_vulnerability(a,'+cmd','manual','EXAMPLE','high')
    s.evaluate_control(a,'NIST-CSF-2.0','PR.AA','partial','fail','Tester','Synthetic failure')
    pdf=export_executive_pdf(s,a,tmp_path/'report.pdf')
    xlsx=export_executive_xlsx(s,a,tmp_path/'report.xlsx')
    assert pdf.read_bytes().startswith(b'%PDF')
    with zipfile.ZipFile(xlsx) as z:assert 'xl/workbook.xml' in z.namelist()
    wb=load_workbook(xlsx)
    assert {'Executive Overview','Risks','Vendors','Vulnerabilities','Control Evaluations','Evidence Inventory','Policies','Incident Inventory'}<=set(wb.sheetnames)
    assert wb['Risks']['A2'].data_type!='f'
    assert wb['Vendors']['A2'].data_type!='f'
    assert wb['Vulnerabilities']['A2'].data_type!='f'


def test_imported_source_name_deduplication_and_invalid_dates(env):
    s,a,b,_=env
    dup=csv_blob(['Incident Name','Status'],[{'Incident Name':'Same','Status':'open'},{'Incident Name':'Same','Status':'closed'}])
    assert s.preview_notion_module(a,dup,'incident_tracker')['skipped']==1
    with pytest.raises(ValidationError):s.import_notion_module(a,dup,'incident_tracker','dup.csv',confirmed=True)
    bad=csv_blob(['Policy Name','Approval Status','Next Review Date'],[{'Policy Name':'Late','Approval Status':'Unknown','Next Review Date':'32/19/2026'}])
    assert s.preview_notion_module(a,bad,'security_policies')['skipped']==1


def test_vault_passphrase_verification(env):
    _s,_a,_b,vault=env
    assert vault.verify_passphrase(PASSWORD)
    assert not vault.verify_passphrase('not-correct')


def test_newer_vault_version_is_not_downgraded(tmp_path):
    path=tmp_path/'future.csvault'
    with Vault(path,PASSWORD,create=True) as vault:
        with vault.transaction() as conn:
            conn.execute('PRAGMA user_version=99')
    with Vault(path,PASSWORD) as vault:
        with pytest.raises(ValidationError, match='newer'):
            Phase3Service(vault)
        assert vault.db.execute('PRAGMA user_version').fetchone()[0]==99


def test_executive_metrics_remediation_without_due_date(env):
    s,a,b,_=env
    risk=s.add_risk(a,'Synthetic issue','',3,3,'Test risk')
    s.add_task(a,risk,'Undated follow-up','Synthetic owner','')
    s.add_task(a,risk,'Past-due follow-up','Synthetic owner',(date.today()-timedelta(days=2)).isoformat())
    assert s.executive_metrics(a)['remediation_overdue']==1
    assert s.executive_metrics(b)['remediation_overdue']==0
