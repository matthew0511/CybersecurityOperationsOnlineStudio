from pathlib import Path
import sqlite3
import zipfile

import pytest
from openpyxl import load_workbook

from cyberstudio.reports import export_pdf, export_xlsx
from cyberstudio.service import StudioService, NotFoundError, ValidationError
from cyberstudio.vault import Vault, VaultAuthenticationError, VaultError, MAGIC

PASS = "test-passphrase-for-demo-2026"


def test_vault_encryption_and_reopen(tmp_path):
    path = tmp_path / "secure.csvault"
    with Vault(path, PASS, create=True) as vault:
        svc = StudioService(vault)
        org = svc.create_organization("Private client A", "Dental")
        assert path.read_bytes().startswith(MAGIC)
        assert b"Private client A" not in path.read_bytes()
        with pytest.raises(VaultError):
            Vault(path, PASS)
    assert not path.read_bytes().startswith(b"SQLite format 3")
    with pytest.raises(VaultAuthenticationError):
        Vault(path, "wrong-passphrase")
    with Vault(path, PASS) as vault:
        assert StudioService(vault).list_organizations()[0]["id"] == org


def test_tamper_detection(tmp_path):
    path = tmp_path / "test.csvault"
    with Vault(path, PASS, create=True) as vault:
        StudioService(vault)
    tampered = bytearray(path.read_bytes())
    tampered[-9] ^= 1
    path.write_bytes(tampered)
    with pytest.raises(VaultAuthenticationError):
        Vault(path, PASS)


def test_no_overwrite_and_short_password(tmp_path):
    path = tmp_path / "test.csvault"
    with pytest.raises(ValueError):
        Vault(path, "short", create=True)
    with Vault(path, PASS, create=True):
        pass
    with pytest.raises(VaultError):
        Vault(path, PASS, create=True)


def test_organization_isolation_and_relationships(tmp_path):
    with Vault(tmp_path / "db.csvault", PASS, create=True) as vault:
        svc = StudioService(vault)
        a = svc.create_organization("Alpha")
        b = svc.create_organization("Bravo")
        ass = svc.create_assessment(a, "Annual assessment")
        svc.answer_question(a, ass, "PR-01", "No", "No MFA observed")
        risk = svc.add_risk(a, "Weak privileged login", "", 4, 5, "Privileged MFA absent", ass)
        task = svc.add_task(a, risk, "Enable MFA", "IT", "2026-11-30")
        assert svc.overview(b) == {"assessments": 0, "risks_open": 0, "risks_high_critical": 0, "tasks_open": 0}
        assert not svc.list_risks(b)
        with pytest.raises(NotFoundError):
            svc.questions(b, ass)
        with pytest.raises(NotFoundError):
            svc.answer_question(b, ass, "PR-01", "Yes", "")
        with pytest.raises(NotFoundError):
            svc.add_risk(b, "Wrong owner", "", 2, 2, "Test", ass)
        with pytest.raises(NotFoundError):
            svc.add_task(b, risk, "Wrong org")
        with pytest.raises(NotFoundError):
            svc.update_task_status(b, task, "done")
        # Database-level composite foreign key prevents accidental cross-org task links.
        with pytest.raises(sqlite3.IntegrityError):
            with vault.transaction() as conn:
                conn.execute("INSERT INTO remediation_tasks VALUES (?,?,?,?,?,?,?,?,?,?)",
                             ("bad", b, risk, "Cross org", "", "", "normal", "open", "", ""))
        svc.update_task_status(a, task, "in_progress", "Work started")
        assert len(svc.task_history(a, task)) == 2
        assert len(svc.audit_events(a)) >= 5
        with pytest.raises(sqlite3.IntegrityError):
            with vault.transaction() as conn:
                conn.execute("DELETE FROM audit_events WHERE organization_id=?", (a,))


@pytest.mark.parametrize("l,i,score,band", [
    (1, 1, 1, "Low"), (2, 2, 4, "Low"), (3, 3, 9, "Moderate"),
    (2, 5, 10, "High"), (4, 4, 16, "High"), (4, 5, 20, "Critical"),
    (5, 5, 25, "Critical"),
])
def test_risk_matrix(l, i, score, band):
    assert StudioService.score(l, i) == (score, band)


@pytest.mark.parametrize("l,i", [(0, 4), (6, 4), (True, 4), (4, 3.1), (-1, 5)])
def test_invalid_risk_matrix(l, i):
    with pytest.raises(ValidationError):
        StudioService.score(l, i)


def test_answers_persist_and_reports(tmp_path):
    path = tmp_path / "save.csvault"
    with Vault(path, PASS, create=True) as vault:
        svc = StudioService(vault)
        org = svc.create_organization("Test Organization")
        ass = svc.create_assessment(org, "Initial baseline")
        assert len(svc.questions(org, ass)) == 8
        svc.answer_question(org, ass, "PR-01", "Partial", "Only some administrators")
        risk = svc.add_risk(org, "Suspicious =SUM(1,1)", "Contains formula", 3, 4, "Manual review")
        svc.add_task(org, risk, "Fix access", "Engineer")
        pdf = export_pdf(svc, org, tmp_path / "risk.pdf")
        xlsx = export_xlsx(svc, org, tmp_path / "risk.xlsx")
        assert pdf.read_bytes().startswith(b"%PDF")
        workbook = load_workbook(xlsx)
        assert workbook.sheetnames == ["Overview", "Risks", "Remediation", "Assessments"]
        assert workbook["Risks"]["A2"].data_type != "f"
        backup = vault.backup(tmp_path / "backup.csvault")
        assert backup.read_bytes() == path.read_bytes()
        with pytest.raises(VaultError):
            vault.backup(backup)
    with Vault(path, PASS) as vault:
        svc = StudioService(vault)
        assert svc.questions(org, ass)[2]["answer"] in {"Partial", "Unanswered"}
        assert next(x for x in svc.questions(org, ass) if x["code"] == "PR-01")["answer"] == "Partial"
        assert len(svc.list_risks(org)) == 1
    with Vault(backup, PASS) as vault:
        assert StudioService(vault).overview(org)["risks_open"] == 1


def test_input_validation(tmp_path):
    with Vault(tmp_path / "db.csvault", PASS, create=True) as vault:
        svc = StudioService(vault)
        with pytest.raises(ValidationError):
            svc.create_organization(" ")
        org = svc.create_organization("One")
        with pytest.raises(ValidationError):
            svc.create_organization("One")
        aid = svc.create_assessment(org, "Risk review")
        with pytest.raises(ValidationError):
            svc.answer_question(org, aid, "PR-01", "Certified", "")
        with pytest.raises(ValidationError):
            svc.add_risk(org, "Missing rationale", "", 3, 4, "")
        rid = svc.add_risk(org, "Valid", "", 2, 5, "Explanation")
        with pytest.raises(ValidationError):
            svc.add_task(org, rid, "Fix", due_date="2026-13-90")
        with pytest.raises(ValidationError):
            svc.update_task_status(org, "invalid", "bogus")
