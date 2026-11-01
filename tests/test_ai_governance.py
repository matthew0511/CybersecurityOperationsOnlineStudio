"""Original-Hub provenance, tenant isolation and configuration-bound release authorization."""
from datetime import date,timedelta
import hashlib,io
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from test_web import web,login,call,PW
from test_experience import invited
from test_platform import post,save
from cyberstudio.web.app import app
from cyberstudio.web.database import connect,service
from cyberstudio.web.platform import ai_governance,catalog

FUTURE=(date.today()+timedelta(days=90)).isoformat()

def setup(web):
    login(web)
    return call(web,'create_organization',name='Synthetic AI Governance'),call(web,'create_organization',name='Isolated AI Tenant')

def case_data(**changes):
    return dict(business_owner='owner',product='Synthetic writing assistant',configuration='Model v1, private pilot baseline',purpose='Public drafting; compare with manual drafting',users='Synthetic staff',prohibited_use='No patient data or autonomous updates',risk_tier='lower',risk_rationale='Public text only; humans verify output',phi='no',clinical='no',patient_facing='no',autonomy='draft_only',data_categories='Public test text only',human_review='Named staff compare against sources and reject unsafe drafts',stop_fallback='Disable feature and use manual drafting',review_on=FUTURE,exit_plan='Revoke access and verify data disposition',**changes)

def values(kind,use_case,**changes):
    result={}
    for key,field in catalog.MODULES[kind][1].items():
        if field['required']:
            result[key]=use_case if key=='use_case_id' else field['options'][0] if field['type']=='select' else 'owner' if key in {'reviewer','response_owner'} else 'Synthetic documented '+key
    result.update(changes);return result

def status(web,org,record,state,expected=200,**extra):
    return post(web,org,'/record/'+record['id']+'/status',dict(version=record['version'],status=state,**extra),expected)

def completed(web,org,record):
    record=status(web,org,record,'in_progress');return status(web,org,record,'completed')

def packet(web,org):
    with connect() as db,db.transaction():
        svc=service(db);ident=svc.add_evidence(org,'Synthetic AI evidence','evaluation.txt',b'Synthetic reviewed evidence only','text/plain')
        svc.review_evidence(org,ident,'verified','owner','Synthetic scoped validation')
        return ident

def ready_case(web,org):
    case=save(web,org,'ai_use_case','Synthetic approved-scope candidate',case_data());eid=packet(web,org)
    save(web,org,'ai_dataflow','Synthetic data map',values('ai_dataflow',case['id'],training_reuse='no'))
    for domain in ai_governance.BASE_REVIEWS:
        completed(web,org,save(web,org,'ai_review',domain,values('ai_review',case['id'],domain=domain,gate='pilot',result='satisfactory',evidence_ids=[eid])))
    for family in ['ordinary_workflow','human_review','resilience']:
        completed(web,org,save(web,org,'ai_evaluation',family,values('ai_evaluation',case['id'],family=family,result='pass',evidence_ids=[eid])))
    return case,eid

def decision(web,org,case,eid,approver='reviewer'):
    return save(web,org,'ai_decision','Scoped pilot decision',values('ai_decision',case['id'],decision_type='release',outcome='restricted',stage='pilot',approver=approver,expires_on=FUTURE,evidence_ids=[eid]))

def test_original_hub_reference_and_register_contract():
    ref=ai_governance.reference()
    assert ref['source_page_id']=='2710be56313383029852016a4e6cb6e7'
    assert ref['parent_page_id']=='df00be5631338303ac1781ee4baff573'
    assert ref['live_sync'] is False and len(ref['source_sha256'])==64
    assert len(ref['sections'])==19
    assert hashlib.sha256(Path('docs/AI_GOVERNANCE_SOURCE_NOTION_2026-10-28.md').read_bytes()).hexdigest()==ref['source_sha256']
    assert '17. Workforce readiness' in ref['sections'][-2]['title']
    assert all(k in catalog.MODULES for k in ai_governance.AI_MODULES)

def test_ai_authentication_scope_and_input_boundaries(web):
    assert web.get('/api/platform/ai/reference').status_code==401
    org,other=setup(web);case=save(web,org,'ai_use_case','Synthetic intake',case_data())
    viewer=invited(web,org,'ai_viewer','viewer')
    assert viewer.get('/api/platform/'+org+'/ai/overview').status_code==200
    assert viewer.get('/api/platform/'+other+'/ai/overview').status_code==403
    post(viewer,org,'/records/ai_use_case',{'title':'Denied','data':case_data()},403)
    post(web,other,'/records/ai_dataflow',{'title':'Cross-tenant link','data':values('ai_dataflow',case['id'])},404)
    for field in ['phi','clinical','patient_facing']:
        data=case_data();data[field]='unknown'
        post(web,org,'/records/ai_use_case',{'title':'Unsafe lower tier','data':data},422)
    data=case_data();data.update(autonomy='consequential_actions',risk_tier='high')
    post(web,org,'/records/ai_use_case',{'title':'Unbounded routine release','data':data},422)
    assert web.get('/api/platform/ai/reference').json()['source_page_id']=='2710be56313383029852016a4e6cb6e7'
    for value in [[],{},True,999]:
        data=case_data();data['phi']=value
        post(web,org,'/records/ai_use_case',{'title':'Malformed field','data':data},422)

def test_ai_readiness_never_implies_approval_and_stale_baselines_fail_closed(web):
    org,_=setup(web);reviewer=invited(web,org,'reviewer','reviewer');case,eid=ready_case(web,org)
    review=web.get('/api/platform/'+org+'/records/ai_review').json()[0]
    post(web,org,'/records/ai_review',{'id':review['id'],'version':review['version'],'title':review['title'],'data':review['data']},409)
    state=web.get('/api/platform/'+org+'/ai/overview').json()['cases'][0]
    assert state['ready_for_human_release_review'] and state['authorization']=='not_authorized'
    record=decision(web,org,case,eid);record=status(web,org,record,'submitted')
    status(web,org,record,'approved',403,password=PW)
    approved=status(reviewer,org,record,'approved',password=PW)
    assert web.get('/api/platform/'+org+'/ai/overview').json()['cases'][0]['authorization']=='pilot'
    updated=case_data();updated['configuration']='Materially changed model v2'
    save(web,org,'ai_use_case',case['title'],updated,id=case['id'],version=case['version'])
    state=web.get('/api/platform/'+org+'/ai/overview').json()['cases'][0]
    assert state['authorization']=='not_authorized' and state['blockers']
    status(web,org,approved,'draft')
    assert web.get('/api/platform/'+org+'/ai/overview').json()['cases'][0]['authorization']=='not_authorized'

def test_ai_missing_evidence_material_events_and_retired_decisions_block_release(web):
    org,_=setup(web);reviewer=invited(web,org,'reviewer','reviewer');case,eid=ready_case(web,org)
    record=decision(web,org,case,eid);record=status(web,org,record,'submitted')
    record=status(reviewer,org,record,'approved',password=PW)
    record=status(web,org,record,'closed')
    assert web.get('/api/platform/'+org+'/ai/overview').json()['cases'][0]['authorization']=='not_authorized'
    event=save(web,org,'ai_event','Unreviewed supplier update',values('ai_event',case['id'],event_type='material_change',severity='high'))
    record=decision(web,org,case,eid);record=status(web,org,record,'submitted')
    status(reviewer,org,record,'approved',409,password=PW)
    status(web,org,event,'closed')
    with connect() as db,db.transaction():service(db).review_evidence(org,eid,'rejected','owner','Evidence no longer supports the tested scope')
    assert not web.get('/api/platform/'+org+'/ai/overview').json()['cases'][0]['ready_for_human_release_review']

def test_ai_frozen_board_reports_are_scoped_and_spreadsheet_safe(web):
    from openpyxl import load_workbook
    org,_=setup(web);save(web,org,'ai_use_case','=1+1',case_data())
    report=post(web,org,'/editions',{'template':'ai_governance','format':'xlsx'})
    response=web.get(report['download_url']);assert response.status_code==200
    book=load_workbook(io.BytesIO(response.content));assert 'Use Cases' in book.sheetnames
    assert all(cell.data_type!='f' for sheet in book for row in sheet for cell in row)
    assert hashlib.sha256(response.content).hexdigest()==report['sha256']
    assert post(web,org,'/editions',{'template':'ai_governance','format':'pdf'})['bytes']>1000
